"""Run exports: one DataFrame, two formats.

CSV via pandas to_csv; XLSX via pandas ExcelWriter with the openpyxl engine.
XLSX cell formatting (see DESIGN.md): emerald header with white bold text,
whole-KHR number format on all money columns (real numeric cells), bold
totals row with top border, auto column widths, frozen header row.
Values are whole KHR, converted to int so Excel treats them as numbers.
"""
import pandas as pd
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TEXT_COLUMNS = ["emp_id", "name", "position", "department"]
MONEY_COLUMNS = [
    "base", "allowance", "overtime", "bonus", "gross", "nssf_employee",
    "taxable", "salary_tax", "advances", "net", "nssf_employer",
]
COLUMNS = TEXT_COLUMNS + MONEY_COLUMNS
LABELS = [
    "Emp ID", "Name", "Position", "Department", "Base", "Allowance",
    "Overtime", "Bonus", "Gross", "NSSF 2%", "Taxable", "Salary tax",
    "Advances", "Net", "Employer NSSF",
]

HEADER_FILL = PatternFill("solid", fgColor="047857")  # DESIGN.md accent-strong
MONEY_FORMAT = "#,##0"


def run_dataframe(run):
    """One row per payslip plus a TOTAL row, labelled columns, int money."""
    rows = []
    for slip in run.payslips.select_related("employee__department", "employee__position"):
        row = {
            "emp_id": slip.employee.emp_id,
            "name": slip.employee.full_name,
            "position": slip.employee.position.name,
            "department": slip.employee.department.name,
        }
        for key in MONEY_COLUMNS:
            row[key] = int(getattr(slip, key))
        rows.append(row)

    totals = {key: sum(r[key] for r in rows) for key in MONEY_COLUMNS}
    totals.update({"emp_id": "", "name": "TOTAL", "position": "", "department": ""})
    rows.append(totals)

    df = pd.DataFrame(rows, columns=COLUMNS)
    df.columns = LABELS
    return df


def run_csv(run) -> str:
    """CSV text; the view adds a UTF-8 BOM so Excel on Windows opens it cleanly."""
    return run_dataframe(run).to_csv(index=False)


def run_xlsx(run) -> bytes:
    """Formatted XLSX bytes (header style, money format, totals, widths)."""
    from io import BytesIO

    buf = BytesIO()
    df = run_dataframe(run)
    sheet = run.period_label[:31]
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=sheet, index=False)
        ws = writer.sheets[sheet]

        # Header: bold white on emerald (5.48:1, AA)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL

        # Column widths from content, capped for readability
        for idx, column in enumerate(df.columns, start=1):
            longest = max([len(str(column))] + [len(str(v)) for v in df[column]])
            ws.column_dimensions[get_column_letter(idx)].width = min(longest + 3, 32)

        # Money columns (5..15) as whole KHR with thousands separators
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.column >= 5:
                    cell.number_format = MONEY_FORMAT

        # Totals row: bold with a top rule
        top = Side(style="thin", color="111827")
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
            cell.border = Border(top=top)

        ws.freeze_panes = "A2"
    return buf.getvalue()
