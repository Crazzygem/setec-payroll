"""Export endpoint tests: auth, CSV/XLSX content, XLSX cell formatting."""
import csv as csv_mod
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from payroll.services import generate_run

from .test_views import make_employee


class ExportTests(TestCase):
    def setUp(self):
        User.objects.create_user("hr", password="hrpass123")
        self.client.login(username="hr", password="hrpass123")
        make_employee()
        self.run, _ = generate_run(2026, 9)
        self.csv_url = reverse("payroll:export_csv", args=[self.run.pk])
        self.xlsx_url = reverse("payroll:export_xlsx", args=[self.run.pk])

    def test_requires_login(self):
        self.client.logout()
        for url in (self.csv_url, self.xlsx_url):
            response = self.client.get(url)
            self.assertEqual(response.status_code, 302)

    def test_csv_content(self):
        response = self.client.get(self.csv_url)
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])
        self.assertIn("payroll_2026-09.csv", response["Content-Disposition"])
        rows = list(csv_mod.reader(response.content.decode("utf-8-sig").splitlines()))
        self.assertEqual(len(rows), 3)          # header + 1 payslip + TOTAL
        self.assertEqual(len(rows[0]), 15)
        self.assertEqual(rows[0][0], "Emp ID")
        slip = self.run.payslips.get()
        self.assertEqual(rows[1][13], str(slip.net))       # Net column
        self.assertEqual(rows[2][1], "TOTAL")
        self.assertEqual(int(rows[2][8]), int(slip.gross))  # Gross totals

    def test_xlsx_content_and_formatting(self):
        response = self.client.get(self.xlsx_url)
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("payroll_2026-09.xlsx", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"PK"))  # valid zip container

        wb = load_workbook(BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ["September 2026"])
        ws = wb.active
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertTrue(ws["A1"].font.bold)
        self.assertEqual(ws["A1"].fill.fgColor.rgb[-6:], "047857")
        self.assertEqual(ws["E2"].number_format, "#,##0")
        slip = self.run.payslips.get()
        self.assertEqual(ws.cell(row=2, column=14).value, int(slip.net))
        last = ws.max_row
        self.assertEqual(ws.cell(row=last, column=2).value, "TOTAL")
        self.assertTrue(ws.cell(row=last, column=2).font.bold)
